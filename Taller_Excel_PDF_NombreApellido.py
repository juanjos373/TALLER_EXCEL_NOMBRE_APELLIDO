from pathlib import Path

try:
    import openpyxl
    from tabulate import tabulate
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
except ImportError as e:
    print(f"Error: {e}. Asegúrese de tener instaladas las librerías necesarias: openpyxl, tabulate, reportlab.")
    exit(1)

BASE_DIR = Path(__file__).resolve().parent
ARCHIVO_EXCEL = BASE_DIR / "ventas_tecnologia_v2.xlsx"
ARCHIVO_PDF = BASE_DIR / "reporte_ejecutivo_ventas.pdf"


def crear_archivo_excel_default(ruta_excel: Path) -> Path:
    if ruta_excel.exists():
        return ruta_excel

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Ventas"
    hoja.append(["ID", "Producto", "Categoría", "Precio ($)", "Ventas (Unid)"])

    datos_base = [
        [1, "Laptop Pro 14", "Computación", 1299.99, 42],
        [2, "Monitor 27\"", "Periféricos", 349.50, 80],
        [3, "Teclado Mecánico", "Accesorios", 89.99, 120],
        [4, "Auriculares Wireless", "Audio", 199.00, 65],
        [5, "Smartphone X", "Celulares", 899.00, 58],
        [6, "Disco SSD 1TB", "Almacenamiento", 129.00, 140],
        [7, "Cámara 4K", "Fotografía", 699.99, 33],
        [8, "Impresora Multifuncional", "Oficina", 249.99, 77],
        [9, "Tablet Pro", "Tablets", 459.00, 52],
        [10, "Router Wi-Fi 6", "Redes", 159.00, 95],
    ]

    for fila in datos_base:
        hoja.append(fila)

    libro.save(ruta_excel)
    return ruta_excel


# 1. Cargar el archivo Excel
crear_archivo_excel_default(ARCHIVO_EXCEL)
wb = openpyxl.load_workbook(ARCHIVO_EXCEL, data_only=True)
hoja = wb.active

# 2. Leer datos dinámicamente
datos = []
for fila in hoja.iter_rows(min_row=2, values_only=True):
    if fila and any(valor is not None for valor in fila):
        datos.append(list(fila))

# 3. Mostrar en consola con tabulate
encabezados = ["ID", "Producto", "Categoría", "Precio ($)", "Ventas (Unid)"]
alineacion = ("center", "left", "center", "right", "right")


def mostrar_tabla(datos_a_mostrar, encabezados_actuales):
    if not datos_a_mostrar:
        print("No hay registros para mostrar.")
        return
    print(tabulate(datos_a_mostrar, headers=encabezados_actuales, tablefmt="fancy_grid", colalign=alineacion))
    print()


def filtrar_por_precio(datos_a_filtrar, precio_minimo=100):
    resultado = []
    for fila in datos_a_filtrar:
        try:
            precio = float(fila[3] or 0)
        except (TypeError, ValueError, IndexError):
            continue
        if precio > precio_minimo:
            resultado.append(fila)
    return resultado


def calcular_metricas(datos_a_calcular):
    total_unidades = sum(int(fila[4] or 0) for fila in datos_a_calcular)
    precio_promedio = sum(float(fila[3] or 0) for fila in datos_a_calcular) / len(datos_a_calcular) if datos_a_calcular else 0
    return total_unidades, precio_promedio


def buscar_por_id(datos_a_buscar, id_buscar):
    resultado = []
    for fila in datos_a_buscar:
        try:
            if int(fila[0]) == id_buscar:
                resultado.append(fila)
        except (TypeError, ValueError, IndexError):
            continue
    return resultado


def generar_pdf(nombre_archivo, datos_a_exportar, encabezados_actuales):
    ruta_pdf = Path(nombre_archivo)
    if ruta_pdf.exists():
        try:
            ruta_pdf.unlink()
        except PermissionError:
            print(f"Advertencia: no se pudo sobrescribir '{ruta_pdf.name}'. Cierre el archivo PDF y vuelva a intentarlo.")
            return

    styles = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle('EstiloTitulo', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=18, textColor=colors.HexColor("#1A365D"), spaceAfter=10)
    estilo_cuerpo = ParagraphStyle('EstiloCuerpo', parent=styles['Normal'], fontSize=9)
    estilo_encabezado = ParagraphStyle('EstiloEncabezado', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, textColor=colors.whitesmoke, alignment=1)

    total_unidades, precio_promedio = calcular_metricas(datos_a_exportar)

    story = [
        Paragraph("Informe Ejecutivo de Ventas Tecno", estilo_titulo),
        Paragraph(f"<b>Total de Productos Procesados:</b> {len(datos_a_exportar)}", estilo_cuerpo),
        Paragraph(f"<b>Total de unidades vendidas:</b> {total_unidades}", estilo_cuerpo),
        Paragraph(f"<b>Precio promedio:</b> ${precio_promedio:,.2f}", estilo_cuerpo),
        Spacer(1, 12)
    ]

    header_paragraphs = [Paragraph(f"<b>{h}</b>", estilo_encabezado) for h in encabezados_actuales]

    filas_pdf = [
        [
            Paragraph(str(r[0]), ParagraphStyle('C', parent=estilo_cuerpo, alignment=1)),
            Paragraph(str(r[1]), ParagraphStyle('L', parent=estilo_cuerpo, alignment=0)),
            Paragraph(str(r[2]), ParagraphStyle('C', parent=estilo_cuerpo, alignment=1)),
            Paragraph(f"${float(r[3]):,.2f}", ParagraphStyle('R', parent=estilo_cuerpo, alignment=2)),
            Paragraph(str(r[4]), ParagraphStyle('R', parent=estilo_cuerpo, alignment=2))
        ]
        for r in datos_a_exportar
    ]

    fila_resumen = [
        Paragraph("Total", ParagraphStyle('C', parent=estilo_cuerpo, alignment=1)),
        Paragraph("Resumen", ParagraphStyle('L', parent=estilo_cuerpo, alignment=0)),
        Paragraph("", ParagraphStyle('C', parent=estilo_cuerpo, alignment=1)),
        Paragraph(f"${precio_promedio:,.2f}", ParagraphStyle('R', parent=estilo_cuerpo, alignment=2)),
        Paragraph(str(total_unidades), ParagraphStyle('R', parent=estilo_cuerpo, alignment=2))
    ]
    filas_pdf.append(fila_resumen)

    anchos_cols = [40, 160, 100, 80, 80]
    tabla_pdf = Table([header_paragraphs] + filas_pdf, colWidths=anchos_cols, repeatRows=1)
    tabla_pdf.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A365D")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFC")]),
        ("BACKGROUND", (0, len(filas_pdf)), (-1, len(filas_pdf)), colors.HexColor("#E2E8F0"))
    ]))

    story.append(tabla_pdf)

    doc = SimpleDocTemplate(str(nombre_archivo), pagesize=letter)
    doc.build(story)


def menu():
    while True:
        print("\n=== MENÚ DE REPORTES ===")
        print("1) Ver catálogo completo")
        print("2) Filtrar por precio (> $100)")
        print("3) Buscar por ID")
        print("4) Exportar informe PDF")
        print("5) Salir")
        opcion = input("Seleccione una opción: ").strip()

        if opcion == "1":
            mostrar_tabla(datos, encabezados)

        elif opcion == "2":
            datos_filtrados = filtrar_por_precio(datos)
            mostrar_tabla(datos_filtrados, encabezados)
            total_unidades, precio_promedio = calcular_metricas(datos_filtrados)
            print(f"Total de unidades vendidas: {total_unidades}")
            print(f"Precio promedio: ${precio_promedio:,.2f}")

        elif opcion == "3":
            try:
                id_buscar = int(input("Ingrese el ID del producto a buscar: "))
            except ValueError:
                print("Debe ingresar un número válido.")
                continue

            resultado = buscar_por_id(datos, id_buscar)
            if resultado:
                mostrar_tabla(resultado, encabezados)
            else:
                print(f"Advertencia: no se encontró el ID {id_buscar}.")

        elif opcion == "4":
            generar_pdf(ARCHIVO_PDF, datos, encabezados)
            print(f"[ÉXITO] Archivo '{ARCHIVO_PDF.name}' generado correctamente.")

        elif opcion == "5":
            print("Saliendo del programa...")
            break

        else:
            print("Opción no válida. Intente nuevamente.")


if __name__ == "__main__":
    menu()